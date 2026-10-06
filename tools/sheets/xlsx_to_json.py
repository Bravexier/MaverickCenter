#!/usr/bin/env python3
"""Convertit un classeur .xlsx (ou un .csv) en JSON uniforme, lu ensuite par les extracteurs de tools/sheets/.

Usage : python3 -I tools/sheets/xlsx_to_json.py <classeur.xlsx|.csv> <sortie.json>
Dépendance : openpyxl (pour les .xlsx seulement) ; le .csv n'a pas de dépendance. Les notes et fils de commentaires sont lus
par tools/sheets/xlsx_comments.py (bibliothèque standard), directement dans le XML du classeur.

Format : { "sheets": [ { "name", "state", "max_row", "max_col", "merged": ["A1:B2", …],
           "validations": [ {"range", "type", "formula1"} ],
           "cells": [ {"r", "c", "v", "f"?, "note"?, "threads"?} ] } ] }
- seules les cellules non vides sont gardées ; r/c = ligne/colonne à partir de 1 ;
- v = valeur en cache (nombres conservés ; dates en ISO 8601), f = formule si la cellule en a une ;
- une plage fusionnée ne porte sa valeur que dans sa cellule en haut à gauche (voir « merged ») ;
- note = la note de la cellule (le petit commentaire jaune), texte tel que le classeur l'écrit : même sens que le champ `note`
  de tools/sheets/sheets_api_to_json.py. Une cellule vide qui porte une note ou un fil est gardée avec v = "" : un commentaire
  n'est jamais perdu en silence (l'extracteur décide quoi en faire ou s'arrête). Sans note, la cellule est inchangée ;
- threads = les fils de commentaires de la cellule (la conversation : le menu « Commentaire » de Google Sheets, distinct de la
  note ; l'API Google Sheets ne les donne pas, d'où l'absence de ce champ dans les dumps de sheets_api_to_json.py), dans l'ordre
  du fichier : [ {"text", "done": true (seulement si le fil est résolu), "replies": [texte, …] (seulement s'il y en a)} ].
  Les réponses « Marquée comme fermée » / « Rouverte » sont des événements d'état écrits par Google, gardés tels quels ;
- une note ou un fil sur une cellule fusionnée autre que celle en haut à gauche reste à SON adresse (openpyxl l'aurait supprimé) ;
- max_row / max_col comptent aussi les cellules qui ne portent qu'un commentaire.
Le contenu du classeur est une donnée non fiable : ce script ne l'exécute jamais.
"""
import csv
import datetime
import importlib.util
import json
import os
import sys
import warnings
import xml.etree.ElementTree as ET
import zipfile


def clean(value):
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    return value


def load_comments_module():
    """tools/sheets/xlsx_comments.py, chargé par son chemin : `python3 -I` n'ajoute pas le dossier du script à sys.path."""
    path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "xlsx_comments.py")
    spec = importlib.util.spec_from_file_location("xlsx_comments", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def attach_comments(cells, data, split_ref):
    """Ajoute `note` et `threads` aux cellules du dump ; une cellule absente (vide) est créée avec v = "". -> (lignes, colonnes)."""
    index = {(x["r"], x["c"]): x for x in cells}
    last_row = last_col = 0
    for ref in set(data["notes"]) | set(data["threads"]):
        pos = split_ref(ref)
        item = index.get(pos)
        if item is None:
            item = index[pos] = {"r": pos[0], "c": pos[1], "v": ""}
            cells.append(item)
        if ref in data["notes"]:
            item["note"] = data["notes"][ref]
        if ref in data["threads"]:
            item["threads"] = data["threads"][ref]
        last_row, last_col = max(last_row, pos[0]), max(last_col, pos[1])
    cells.sort(key=lambda x: (x["r"], x["c"]))  # ordre des lignes puis des colonnes, comme iter_rows
    return last_row, last_col


def from_xlsx(path):
    import openpyxl  # importé ici : inutile pour un .csv

    comments_module = load_comments_module()
    try:
        comments = comments_module.read_comments(path)
    except (ValueError, KeyError, zipfile.BadZipFile, ET.ParseError) as err:
        sys.exit("%s : lecture des commentaires impossible : %s" % (path, err))
    with warnings.catch_warnings():
        # openpyxl supprime la note d'une cellule fusionnée non principale en le signalant ; elle est relue dans le XML.
        warnings.filterwarnings("ignore", message=".*is part of a merged range but has a comment.*")
        values = openpyxl.load_workbook(path, data_only=True)
        formulas = openpyxl.load_workbook(path, data_only=False)
    sheets = []
    for ws in values.worksheets:
        wf = formulas[ws.title]
        cells = []
        for row in ws.iter_rows():
            for cell in row:
                formula = wf[cell.coordinate].value
                is_formula = isinstance(formula, str) and formula.startswith("=")
                if cell.value is None and not is_formula:
                    continue
                item = {"r": cell.row, "c": cell.column, "v": clean(cell.value)}
                if is_formula:
                    item["f"] = formula
                cells.append(item)
        max_row, max_col = ws.max_row, ws.max_column
        data = comments.pop(ws.title, None)
        if data:
            last_row, last_col = attach_comments(cells, data, comments_module.split_ref)
            max_row, max_col = max(max_row, last_row), max(max_col, last_col)
        validations = []
        for dv in ws.data_validations.dataValidation:
            validations.append({"range": str(dv.sqref), "type": dv.type, "formula1": dv.formula1})
        sheets.append({
            "name": ws.title,
            "state": ws.sheet_state,
            "max_row": max_row,
            "max_col": max_col,
            "merged": [str(r) for r in ws.merged_cells.ranges],
            "validations": validations,
            "cells": cells,
        })
    lost = {name: sorted(set(d["notes"]) | set(d["threads"])) for name, d in comments.items() if d["notes"] or d["threads"]}
    if lost:  # feuille de graphique ou autre : aucun commentaire ne se perd en silence
        sys.exit("commentaires dans des onglets absents du dump (pas des feuilles de calcul) : %s" % lost)
    return {"sheets": sheets}


def from_csv(path):
    cells = []
    max_col = 0
    with open(path, newline="", encoding="utf-8") as handle:
        rows = list(csv.reader(handle))
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            if value != "":
                cells.append({"r": r, "c": c, "v": value})
                max_col = max(max_col, c)
    sheet = {"name": "Feuille 1", "state": "visible", "max_row": len(rows), "max_col": max_col,
             "merged": [], "validations": [], "cells": cells}
    return {"sheets": [sheet]}


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    source, target = sys.argv[1], sys.argv[2]
    data = from_csv(source) if source.lower().endswith(".csv") else from_xlsx(source)
    with open(target, "w", encoding="utf-8") as handle:
        json.dump(data, handle, ensure_ascii=False, separators=(",", ":"))
    total = sum(len(s["cells"]) for s in data["sheets"])
    print(f"{len(data['sheets'])} onglet(s), {total} cellules -> {target}")
    cells = [c for s in data["sheets"] for c in s["cells"]]
    notes = sum(1 for c in cells if "note" in c)
    threads = [t for c in cells for t in c.get("threads", [])]
    if notes or threads:
        replies = sum(len(t.get("replies", [])) for t in threads)
        carrying = sum(1 for c in cells if "note" in c or "threads" in c)
        print(f"  commentaires de cellule : {carrying} cellule(s) en portent : {notes} note(s), "
              f"{len(threads)} fil(s) de commentaires ({replies} réponse(s))")


if __name__ == "__main__":
    main()
